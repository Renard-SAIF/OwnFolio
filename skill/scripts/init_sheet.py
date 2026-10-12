#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""建一个标准表头的空项目库。

用法:
    python init_sheet.py --out "文档/项目库.xlsx"     # 生成 xlsx（需要 openpyxl）
    python init_sheet.py --out "文档/项目库.csv"      # 生成两份 csv（零依赖兜底）

xlsx 会建两张工作表：项目库 + 看法流水，首行冻结、列宽调好、表头加底色。
"""

import argparse
import csv
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(os.path.dirname(HERE), 'assets')

MAIN = [
    ('company', '公司名称', 18), ('industry', '行业标签', 14),
    ('round', '融资轮次', 12), ('status', '决策状态', 14),
    ('sourceType', '项目来源', 14), ('sourceDetail', '来源明细', 16),
    ('score', '兴趣度', 8), ('thesis', '一句话投资逻辑', 42),
    ('meetDate', '交流日期', 12), ('attendees', '对方参会人及角色', 22),
    ('product', '产品与技术', 32), ('competitors', '竞争对手与差异点', 28),
    ('revenue', '收入与财务', 20), ('financing', '本轮融资与估值', 26),
    ('credibility', '信息可信度', 12), ('myView', '我的看法', 42),
    ('concernResult', '当初的预测应验了吗', 16), ('missReason', '为什么没中', 26),
    ('dropReason', '放弃原因', 22), ('nextFollow', '下次跟进日期', 14),
    ('createdAt', '建档时间', 12), ('updatedAt', '最后更新', 12),
]

LOG = [
    ('company', '公司名称', 18), ('at', '时点', 12),
    ('stage', '当时阶段', 14), ('score', '当时兴趣度', 12),
    ('view', '当时的看法', 46), ('expect', '当时的预测', 34),
    ('result', '后来的结果', 14), ('missReason', '为什么没中', 26),
    ('source', '来源', 18),
]


def write_xlsx(path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        print('没有 openpyxl，改成输出两份 csv')
        return False

    wb = Workbook()
    head_font = Font(bold=True, color='1F2937', size=11)
    head_fill = PatternFill('solid', fgColor='EEF2F7')

    for idx, cols in enumerate((MAIN, LOG)):
        ws = wb.active if idx == 0 else wb.create_sheet()
        ws.title = '项目库' if idx == 0 else '看法流水'
        for c, (_, label, width) in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=c, value=label)
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(vertical='center')
            ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = width
        ws.freeze_panes = 'A2'
        ws.row_dimensions[1].height = 24

    wb.save(path)
    return True


def write_csv(path):
    base = path[:-4] if path.lower().endswith('.csv') else path
    for name, cols in (('项目库', MAIN), ('看法流水', LOG)):
        out = '%s-%s.csv' % (base, name)
        with open(out, 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.writer(f)
            w.writerow([c[0] for c in cols])
            w.writerow([c[1] for c in cols])
        print('已生成 %s' % out)
    return True


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', required=True, help='输出路径，.xlsx 或 .csv')
    args = ap.parse_args()

    out = os.path.abspath(args.out)
    parent = os.path.dirname(out)
    if parent and not os.path.isdir(parent):
        try:
            os.makedirs(parent, exist_ok=True)
        except OSError as e:
            print('建不了目录：%s' % e)
            return 1

    ok = False
    if out.lower().endswith('.xlsx'):
        ok = write_xlsx(out)
        if not ok:
            write_csv(out)
    elif out.lower().endswith('.csv'):
        ok = write_csv(out)
    else:
        print('只认 .xlsx 或 .csv')
        return 1

    if ok and out.lower().endswith('.xlsx'):
        print('已生成 %s（两张工作表：项目库 / 看法流水）' % out)
    return 0


if __name__ == '__main__':
    sys.exit(main())
